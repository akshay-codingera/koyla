import os
import sys
import time
import tracemalloc
from pathlib import Path

import openpyxl
import pytest

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.services.parsers.spreadsheet_parser import spreadsheet_parser
from app.services.parsers.csv_parser import csv_parser


def test_normal_xlsx(tmp_path):
    """Verify standard single-sheet XLSX parsing and table extraction."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Monthly_Production"
    ws.append(["Month", "Production_MT", "Target_MT", "Grade"])
    ws.append(["April 2024", "45.2", "42.0", "G11"])
    ws.append(["May 2024", "48.1", "45.0", "G11"])
    ws.append(["June 2024", "43.7", "44.0", "G12"])

    file_path = str(tmp_path / "normal_prod.xlsx")
    wb.save(file_path)
    wb.close()

    parsed = spreadsheet_parser.parse(file_path)
    assert parsed.total_pages == 1
    assert len(parsed.tables) == 1
    table = parsed.tables[0]
    assert table.headers == ["Month", "Production_MT", "Target_MT", "Grade"]
    assert len(table.rows) == 3
    assert table.rows[0][0] == "April 2024"
    assert table.rows[0][1] == "45.2"


def test_multi_sheet_xlsx(tmp_path):
    """Verify multi-sheet XLSX creates distinct pages and tables for each sheet."""
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Seam_Reserves"
    ws1.append(["Seam", "Thickness_M", "Reserve_MT"])
    ws1.append(["Seam I", "4.2", "120.5"])

    ws2 = wb.create_sheet("Quality_Parameters")
    ws2.append(["Parameter", "Value", "Unit"])
    ws2.append(["Ash", "34.5", "%"])
    ws2.append(["GCV", "4120", "kcal/kg"])

    ws3 = wb.create_sheet("Empty_Sheet")

    file_path = str(tmp_path / "multi_sheet.xlsx")
    wb.save(file_path)
    wb.close()

    parsed = spreadsheet_parser.parse(file_path)
    assert parsed.total_pages == 3
    assert len(parsed.tables) == 2  # 2 sheets with data, 1 empty
    assert parsed.metadata["sheet_count"] == 3

    sheet1_table = next(t for t in parsed.tables if "Seam_Reserves" in t.caption)
    assert sheet1_table.rows[0][0] == "Seam I"

    sheet2_table = next(t for t in parsed.tables if "Quality_Parameters" in t.caption)
    assert sheet2_table.rows[1][0] == "GCV"


def test_large_synthetic_xlsx_streaming(tmp_path):
    """Verify large synthetic XLSX (5,000 rows) processes with bounded memory and chunking."""
    file_path = str(tmp_path / "large_5000_rows.xlsx")

    # Generate 5,000 rows
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dispatch_Register"
    ws.append(["Challan_No", "Truck_No", "Weight_Tonnes", "Grade", "Destination"])

    for i in range(1, 5001):
        ws.append([f"CH-{i:06d}", f"JH-01-AB-{i%9000+1000}", f"{30 + (i % 15):.2f}", "G11", "Patratu TPS"])

    wb.save(file_path)
    wb.close()

    # Measure memory and latency during parsing
    tracemalloc.start()
    t0 = time.perf_counter()

    parsed = spreadsheet_parser.parse(file_path, chunk_size=1500)

    elapsed = time.perf_counter() - t0
    current_mem, peak_mem = tracemalloc.get_traced_memory()
    tracemalloc.stop()

    peak_mb = peak_mem / (1024 * 1024)
    print(f"\n[XLSX 5,000 Rows Benchmark] Elapsed: {elapsed:.2f}s | Peak RAM: {peak_mb:.2f} MB")

    # Assertions
    assert parsed.total_pages == 1
    # 5,000 rows with chunk_size=1500 -> 4 chunks (1500, 1500, 1500, 500)
    assert len(parsed.tables) == 4
    total_rows = sum(len(t.rows) for t in parsed.tables)
    assert total_rows == 5000
    assert parsed.tables[0].part_number == 1
    assert parsed.tables[1].is_continuation is True
    assert parsed.tables[1].part_number == 2
    assert elapsed < 15.0  # reasonable performance ceiling


def test_large_csv_streaming(tmp_path):
    """Verify large CSV (10,000 rows) processes with streaming reader and chunking."""
    file_path = str(tmp_path / "large_10000_rows.csv")

    with open(file_path, "w", encoding="utf-8") as f:
        f.write("Borehole_ID,Depth_M,Strata_Type,Core_Recovery_Pct\n")
        for i in range(1, 10001):
            f.write(f"BH-2024-{i:05d},{i * 0.5:.2f},SANDSTONE_MEDIUM_GRAINED,{85 + (i % 15)}\n")

    tracemalloc.start()
    t0 = time.perf_counter()

    parsed = csv_parser.parse(file_path, chunk_size=2500)

    elapsed = time.perf_counter() - t0
    current_mem, peak_mem = tracemalloc.get_traced_memory()
    tracemalloc.stop()

    peak_mb = peak_mem / (1024 * 1024)
    print(f"\n[CSV 10,000 Rows Benchmark] Elapsed: {elapsed:.2f}s | Peak RAM: {peak_mb:.2f} MB")

    assert parsed.total_pages == 1
    # 10,000 rows with chunk_size=2500 -> 4 chunks
    assert len(parsed.tables) == 4
    total_rows = sum(len(t.rows) for t in parsed.tables)
    assert total_rows == 10000
    assert parsed.tables[0].headers == ["Borehole_ID", "Depth_M", "Strata_Type", "Core_Recovery_Pct"]
    assert elapsed < 5.0  # fast streaming processing


def test_malformed_spreadsheet_handling(tmp_path):
    """Verify malformed/corrupted spreadsheet does not crash unhandled."""
    corrupt_file = str(tmp_path / "corrupt.xlsx")
    with open(corrupt_file, "wb") as f:
        f.write(b"NOT_A_VALID_ZIP_OR_XLSX_STREAM_CORRUPTED_BYTES_HERE")

    with pytest.raises(Exception):
        spreadsheet_parser.parse(corrupt_file)
