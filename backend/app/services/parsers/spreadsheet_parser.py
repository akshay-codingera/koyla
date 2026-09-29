import logging
import uuid
from pathlib import Path
from typing import List, Any, Optional
import openpyxl
from app.services.parsers.base import BaseParser, ParsedDocument, ParsedPage, ParsedTable
from app.services.parsers.spreadsheet_headers import (
    extract_merged_ranges_from_worksheet,
    detect_header_depth,
    propagate_hierarchical_headers,
)

logger = logging.getLogger("app.services.parsers.spreadsheet_parser")


class SpreadsheetParser(BaseParser):
    """
    Hardened Spreadsheet Parser for Excel (.xlsx) workbooks.
    Features:
    - Multi-row merged-header detection and hierarchical propagation (Phase 11 P1-2)
    - Memory-conscious read_only streaming mode with zip-based merge range fallback
    - Incremental sheet iteration without full DOM instantiation
    - Configurable chunk size for very large sheets
    - Table provenance and hierarchical header metadata preservation
    - Fallback to standard loader if read_only fails
    """
    DEFAULT_CHUNK_SIZE = 2000

    def can_handle(self, mime_type: str, filename: str) -> bool:
        ext = Path(filename).suffix.lower()
        if ext == ".xls":
            return False
        return ext == ".xlsx" or "openxmlformats-officedocument.spreadsheetml" in mime_type

    def parse(self, file_path: str, chunk_size: Optional[int] = None) -> ParsedDocument:
        chunk_size = chunk_size or self.DEFAULT_CHUNK_SIZE
        
        # Attempt memory-efficient read_only streaming first
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            read_only_mode = True
        except Exception:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            read_only_mode = False

        pages: List[ParsedPage] = []
        all_tables: List[ParsedTable] = []

        try:
            for s_idx, sheet_name in enumerate(wb.sheetnames):
                sheet = wb[sheet_name]
                page_num = s_idx + 1

                # 1. Extract merged ranges for this worksheet
                merged_ranges = extract_merged_ranges_from_worksheet(
                    ws=sheet,
                    file_path=file_path,
                    sheet_name=sheet_name,
                    sheet_idx=s_idx + 1
                )

                headers: Optional[List[str]] = None
                raw_headers: Optional[List[str]] = None
                header_depth: int = 1
                current_chunk_rows: List[List[Any]] = []
                sheet_tables: List[ParsedTable] = []
                total_data_rows = 0
                first_100_rows: List[List[str]] = []
                sheet_logical_id = str(uuid.uuid4())
                part_idx = 0

                row_iter = sheet.iter_rows(values_only=True)

                # 2. Peek initial non-empty rows to evaluate multi-row header structure
                initial_non_empty_rows: List[List[Any]] = []
                for row in row_iter:
                    if any(cell is not None for cell in row):
                        initial_non_empty_rows.append(list(row))
                        if len(initial_non_empty_rows) >= 12:
                            break

                if not initial_non_empty_rows:
                    # Empty sheet
                    pages.append(
                        ParsedPage(
                            page_number=page_num,
                            text=f"Sheet: {sheet_name} (Empty)",
                            ocr_applied=False,
                            confidence=1.0,
                            tables=[],
                            metadata={
                                "sheet_name": sheet_name,
                                "empty": True,
                                "read_only_mode": read_only_mode
                            }
                        )
                    )
                    continue

                # 3. Detect header depth and propagate hierarchical headers
                header_depth = detect_header_depth(initial_non_empty_rows, merged_ranges)
                header_slice = initial_non_empty_rows[:header_depth]
                data_slice = initial_non_empty_rows[header_depth:]

                headers, raw_headers = propagate_hierarchical_headers(
                    header_rows=header_slice,
                    merged_ranges=merged_ranges
                )

                logger.info(
                    f"Parsed sheet '{sheet_name}': detected header depth {header_depth}, "
                    f"{len(merged_ranges)} merged ranges. Columns: {headers}"
                )

                def append_data_row(row_cells: List[Any]):
                    nonlocal total_data_rows, part_idx, current_chunk_rows, sheet_tables, all_tables
                    cleaned_row = [str(c).strip() if c is not None else "" for c in row_cells]
                    # Normalize length against headers
                    if len(cleaned_row) < len(headers):
                        cleaned_row = cleaned_row + [""] * (len(headers) - len(cleaned_row))
                    elif len(cleaned_row) > len(headers):
                        cleaned_row = cleaned_row[:len(headers)]

                    current_chunk_rows.append(cleaned_row)
                    total_data_rows += 1

                    if len(first_100_rows) < 100:
                        first_100_rows.append(cleaned_row)

                    # Chunking for very large sheets
                    if len(current_chunk_rows) >= chunk_size:
                        part_idx += 1
                        caption = f"Sheet: {sheet_name}" if part_idx == 1 else f"Sheet: {sheet_name} (Part {part_idx})"
                        table_part = ParsedTable(
                            page_number=page_num,
                            table_index=part_idx,
                            caption=caption,
                            headers=headers,
                            rows=current_chunk_rows,
                            logical_table_id=sheet_logical_id,
                            is_continuation=part_idx > 1,
                            part_number=part_idx,
                            metadata={
                                "sheet_name": sheet_name,
                                "row_count": len(current_chunk_rows),
                                "col_count": len(headers),
                                "part_number": part_idx,
                                "logical_table_id": sheet_logical_id,
                                "is_continuation": part_idx > 1,
                                "read_only_mode": read_only_mode,
                                "header_depth": header_depth,
                                "has_hierarchical_headers": header_depth > 1,
                                "hierarchical_headers": headers,
                                "raw_headers": raw_headers,
                                "merged_ranges_count": len(merged_ranges),
                                "merged_ranges": [str(m) for m in merged_ranges]
                            }
                        )
                        sheet_tables.append(table_part)
                        all_tables.append(table_part)
                        current_chunk_rows = []

                # 4. Stream initial data slice
                for r in data_slice:
                    append_data_row(r)

                # 5. Stream remaining rows from sheet iterator
                for row in row_iter:
                    if not any(cell is not None for cell in row):
                        continue
                    append_data_row(list(row))

                # Add final / remaining chunk
                if current_chunk_rows or not sheet_tables:
                    part_idx += 1
                    caption = f"Sheet: {sheet_name}" if part_idx == 1 else f"Sheet: {sheet_name} (Part {part_idx})"
                    table_part = ParsedTable(
                        page_number=page_num,
                        table_index=part_idx,
                        caption=caption,
                        headers=headers,
                        rows=current_chunk_rows,
                        logical_table_id=sheet_logical_id,
                        is_continuation=part_idx > 1,
                        part_number=part_idx,
                        metadata={
                            "sheet_name": sheet_name,
                            "row_count": len(current_chunk_rows),
                            "col_count": len(headers),
                            "part_number": part_idx,
                            "logical_table_id": sheet_logical_id,
                            "is_continuation": part_idx > 1,
                            "read_only_mode": read_only_mode,
                            "header_depth": header_depth,
                            "has_hierarchical_headers": header_depth > 1,
                            "hierarchical_headers": headers,
                            "raw_headers": raw_headers,
                            "merged_ranges_count": len(merged_ranges),
                            "merged_ranges": [str(m) for m in merged_ranges]
                        }
                    )
                    sheet_tables.append(table_part)
                    all_tables.append(table_part)

                # Generate structured sheet text summary for search index
                text_lines = [f"# Sheet: {sheet_name}", f"Headers: {' | '.join(headers)}", "Data:"]
                for r in first_100_rows:
                    text_lines.append(" | ".join(r))
                if total_data_rows > 100:
                    text_lines.append(f"... ({total_data_rows - 100} more rows)")

                page_text = "\n".join(text_lines)

                pages.append(
                    ParsedPage(
                        page_number=page_num,
                        text=page_text,
                        ocr_applied=False,
                        confidence=1.0,
                        tables=sheet_tables,
                        metadata={
                            "sheet_name": sheet_name,
                            "row_count": total_data_rows,
                            "parts": part_idx,
                            "read_only_mode": read_only_mode,
                            "header_depth": header_depth,
                            "has_hierarchical_headers": header_depth > 1
                        }
                    )
                )

            return ParsedDocument(
                pages=pages,
                total_pages=len(pages),
                tables=all_tables,
                metadata={
                    "format": "Spreadsheet",
                    "sheet_count": len(wb.sheetnames),
                    "sheet_names": wb.sheetnames,
                    "read_only_mode": read_only_mode
                },
                ocr_applied=False
            )
        finally:
            wb.close()


spreadsheet_parser = SpreadsheetParser()
