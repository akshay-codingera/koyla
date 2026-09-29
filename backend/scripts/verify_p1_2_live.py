"""
Koyla Phase 11 P1-2: Multi-Row Merged-Header Propagation Live Verification

Validates end-to-end container-level workflow:
1. Generates synthetic multi-row merged-header XLSX workbook:
                Production              Quality
   Mine         Raw Coal   Washed Coal  Ash %   Moisture %
   Mine-A       1200       850          18.2    4.1
   Mine-B       1300       900          17.4    4.5
2. Ingests workbook through real parser pipeline
3. Verifies merged-range detection and hierarchical header propagation:
   ['Mine', 'Production | Raw Coal', 'Production | Washed Coal', 'Quality | Ash %', 'Quality | Moisture %']
4. Verifies database persistence in PostgreSQL (Table, TableRow, ExtractedField)
5. Verifies provenance (Sheet, cell coordinates B3 and B4)
6. Executes Grounded Q&A / Structured Numerical Aggregation query:
   "What was the total raw coal production?"
   Validates calculated deterministic total: 1200 + 1300 = 2500 MT
"""

import os
import sys
import uuid
import openpyxl
from pathlib import Path

# Add app to path
sys.path.insert(0, "/app")

from app.db.database import SessionLocal
from app.models.organization import Organization
from app.models.user import User, UserRole
from app.models.document import Document, DocumentPage, Table, TableRow
from app.models.extraction import ExtractedField, ExtractionRun
from app.services.parsers.spreadsheet_parser import spreadsheet_parser
from app.services.extraction.rule_based import RuleBasedExtractionProvider
from app.services.qa.structured_lookup import structured_lookup_service
from app.services.qa.qa_service import qa_service
from app.core.security import create_access_token


def run_live_verification():
    print("=" * 70)
    print("KOYLA PHASE 11 P1-2: HIERARCHICAL SPREADSHEET HEADER VERIFICATION")
    print("=" * 70)

    db = SessionLocal()
    xlsx_path = "/tmp/koyla_live_hierarchical_prod.xlsx"

    try:
        # 1. Setup Test Organization and User (clean up prior runs if present)
        existing_org = db.query(Organization).filter(Organization.code == "P12_LIVE_SUBSIDIARY").first()
        if existing_org:
            db.query(ExtractedField).filter(ExtractedField.organization_id == existing_org.id).delete(synchronize_session=False)
            db.query(Document).filter(Document.organization_id == existing_org.id).delete(synchronize_session=False)
            user_ids = [u.id for u in db.query(User).filter(User.organization_id == existing_org.id).all()]
            if user_ids:
                db.query(UserRole).filter(UserRole.user_id.in_(user_ids)).delete(synchronize_session=False)
            db.query(User).filter(User.organization_id == existing_org.id).delete(synchronize_session=False)
            db.delete(existing_org)
            db.commit()

        org = Organization(
            name="CMPDI_LIVE_P12",
            code="P12_LIVE_SUBSIDIARY",
            org_type="SUBSIDIARY"
        )
        db.add(org)
        db.commit()
        db.refresh(org)

        user = User(
            username=f"geologist_p12_{uuid.uuid4().hex[:6]}",
            hashed_password="hashed_pw_dummy",
            full_name="P12 Verification Officer",
            email="p12_officer@cmpdi.gov.in",
            organization_id=org.id,
            is_active=True
        )
        db.add(user)
        db.commit()
        db.refresh(user)

        from app.models.user import Role
        role = db.query(Role).filter(Role.code == "MINING_OFFICER").first()
        if not role:
            role = db.query(Role).first()
        if role:
            user_role = UserRole(user_id=user.id, role_id=role.id)
            db.add(user_role)
            db.commit()

        print(f"[*] Created organization '{org.name}' ({org.id}) and user '{user.username}'")

        # 2. Generate Synthetic XLSX Workbook with Hierarchical Merged Headers
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Monthly_Production"

        # Row 1: Merged super-headers
        ws.merge_cells("B1:C1")
        ws["B1"] = "Production"
        ws.merge_cells("D1:E1")
        ws["D1"] = "Quality"

        # Row 2: Sub-headers
        ws.append(["Mine", "Raw Coal", "Washed Coal", "Ash %", "Moisture %"])

        # Row 3 & 4: Data rows
        ws.append(["Mine-A", 1200, 850, 18.2, 4.1])
        ws.append(["Mine-B", 1300, 900, 17.4, 4.5])

        wb.save(xlsx_path)
        wb.close()
        print(f"[*] Generated synthetic workbook at '{xlsx_path}' with 2-row merged headers")

        # 3. Parse via SpreadsheetParser
        parsed_doc = spreadsheet_parser.parse(xlsx_path)
        assert parsed_doc.total_pages == 1
        assert len(parsed_doc.tables) == 1
        tbl = parsed_doc.tables[0]

        expected_headers = [
            "Mine",
            "Production | Raw Coal",
            "Production | Washed Coal",
            "Quality | Ash %",
            "Quality | Moisture %"
        ]
        assert tbl.headers == expected_headers
        assert tbl.metadata.get("has_hierarchical_headers") is True
        assert tbl.metadata.get("header_depth") == 2
        assert len(tbl.rows) == 2
        assert tbl.rows[0] == ["Mine-A", "1200", "850", "18.2", "4.1"]
        assert tbl.rows[1] == ["Mine-B", "1300", "900", "17.4", "4.5"]

        print("[+] SpreadsheetParser successfully propagated merged headers:")
        for idx, h in enumerate(tbl.headers, start=1):
            print(f"    Col {idx}: {h}")

        # 4. Persist to PostgreSQL (Document, DocumentPage, Table, TableRow)
        doc = Document(
            organization_id=org.id,
            title="Live Hierarchical Coal Production Report",
            original_filename="live_hierarchical_prod.xlsx",
            file_path=xlsx_path,
            sha256_hash="sha256_live_p12_hier",
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            document_type="PRODUCTION_REPORT",
            file_size_bytes=os.path.getsize(xlsx_path)
        )
        db.add(doc)
        db.commit()
        db.refresh(doc)

        page = DocumentPage(
            document_id=doc.id,
            page_number=1,
            extracted_text=parsed_doc.pages[0].text
        )
        db.add(page)
        db.commit()
        db.refresh(page)

        db_table = Table(
            document_id=doc.id,
            page_number=1,
            table_index=1,
            caption="Monthly_Production",
            headers=tbl.headers,
            row_count=len(tbl.rows),
            col_count=len(tbl.headers),
            metadata_json=tbl.metadata
        )
        db.add(db_table)
        db.commit()
        db.refresh(db_table)

        row1 = TableRow(
            table_id=db_table.id,
            row_index=1,
            source_page=1,
            cells=tbl.rows[0]
        )
        row2 = TableRow(
            table_id=db_table.id,
            row_index=2,
            source_page=1,
            cells=tbl.rows[1]
        )
        db.add_all([row1, row2])
        db.commit()
        db_table.table_rows = [row1, row2]

        print(f"[+] Persisted Document ({doc.id}) and Table with {len(db_table.table_rows)} TableRows to PostgreSQL")

        # 5. Extract Structured Fields via Rule-Based Provider
        extractor = RuleBasedExtractionProvider()
        candidates = extractor.extract(chunks=[], tables=[db_table], pages=[page])

        run = ExtractionRun(
            document_id=doc.id,
            status="COMPLETED",
            provider="RULE_BASED"
        )
        db.add(run)
        db.commit()
        db.refresh(run)

        persisted_fields = []
        for cand in candidates:
            ef = ExtractedField(
                extraction_run_id=run.id,
                organization_id=org.id,
                document_id=doc.id,
                field_name=cand.field_name,
                field_category=cand.field_category,
                data_type=cand.data_type,
                raw_value=cand.raw_value,
                normalized_value=cand.normalized_value,
                numeric_value=cand.numeric_value,
                unit=cand.unit,
                page_number=cand.page_number,
                table_id=db_table.id,
                source_text=cand.source_text,
                extraction_method=cand.extraction_method,
                confidence_score=cand.confidence_score,
                confidence_level=cand.confidence_level,
                verification_status="VERIFIED",
                metadata_json=cand.metadata
            )
            db.add(ef)
            persisted_fields.append(ef)
        db.commit()

        prod_fields = [f for f in persisted_fields if f.field_name == "production_quantity"]
        assert len(prod_fields) == 2, f"Expected 2 production_quantity fields, got {len(prod_fields)}"
        assert prod_fields[0].numeric_value == 1200.0
        assert prod_fields[1].numeric_value == 1300.0
        assert prod_fields[0].metadata_json.get("cell_coordinate") == "B3"
        assert prod_fields[1].metadata_json.get("cell_coordinate") == "B4"
        assert prod_fields[0].metadata_json.get("header") == "Production | Raw Coal"

        print("[+] Extracted structured domain fields from hierarchical columns:")
        for pf in prod_fields:
            coord = pf.metadata_json.get("cell_coordinate")
            print(f"    - {pf.metadata_json.get('entity_name')}: {pf.numeric_value} {pf.unit} [Header: '{pf.metadata_json.get('header')}', Coordinate: {coord}]")

        # 6. Execute P0-3 Structured Numerical Aggregation Query
        query = "What was the total raw coal production?"
        print(f"[*] Executing structured aggregation query: \"{query}\"")

        lookup_res = structured_lookup_service.lookup(
            db=db,
            query=query,
            allowed_org_ids=[org.id]
        )

        assert lookup_res.is_aggregation_query is True, "Query was not recognized as aggregation"
        assert lookup_res.aggregation_result is not None, "AggregationResult was None"
        agg = lookup_res.aggregation_result

        print(f"[+] Multi-Document Numerical Aggregation Result:")
        print(f"    - Metric:           {agg.metric_name}")
        print(f"    - Operation:        {agg.operation}")
        print(f"    - Calculated Total: {agg.calculated_value} {agg.unit}")
        print(f"    - Record Count:     {agg.record_count}")
        print(f"    - Formula:          {agg.formula}")
        print(f"    - Summary:          {agg.natural_language_summary}")

        assert agg.calculated_value == 2500.0, f"Expected 2500.0, got {agg.calculated_value}"
        assert agg.record_count == 2
        assert agg.unit == "MT"
        assert any(c.entity_name == "Mine-A" for c in agg.contributing_records)
        assert any(c.entity_name == "Mine-B" for c in agg.contributing_records)

        # 7. Execute Grounded QA Service Query
        qa_res = qa_service.answer_query(
            db=db,
            current_user=user,
            query=query,
            allowed_org_ids=[org.id]
        )
        print(f"[+] Grounded QA Service Output:")
        print(f"    - Verification Status: {qa_res['verification_status']}")
        print(f"    - Answer Snippet:      \"{qa_res['answer'][:140]}...\"")
        print(f"    - Citations:           {len(qa_res['citations'])}")

        assert qa_res["verification_status"] in ("VERIFIED", "SUPPORTED", "PARTIALLY_SUPPORTED")
        assert len(qa_res["citations"]) >= 1

        print("=" * 70)
        print("ALL P1-2 VERIFICATION CHECKS PASSED SUCCESSFULLY (100% GREEN)")
        print("=" * 70)

        # Cleanup
        db.query(ExtractedField).filter(ExtractedField.organization_id == org.id).delete()
        db.query(ExtractionRun).filter(ExtractionRun.document_id == doc.id).delete()
        db.query(TableRow).filter(TableRow.table_id == db_table.id).delete()
        db.query(Table).filter(Table.document_id == doc.id).delete()
        db.query(DocumentPage).filter(DocumentPage.document_id == doc.id).delete()
        db.query(Document).filter(Document.id == doc.id).delete()
        db.query(UserRole).filter(UserRole.user_id == user.id).delete()
        db.query(User).filter(User.id == user.id).delete()
        db.query(Organization).filter(Organization.id == org.id).delete()
        db.commit()

        if os.path.exists(xlsx_path):
            os.remove(xlsx_path)

    finally:
        db.close()


if __name__ == "__main__":
    run_live_verification()
