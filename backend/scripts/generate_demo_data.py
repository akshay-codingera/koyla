import os
import zipfile
import pymupdf
import docx
import openpyxl
from PIL import Image, ImageDraw

def generate_demo_files():
    out_dir = "demo_data"
    os.makedirs(out_dir, exist_ok=True)
    
    # 1. Clean Digital PDF
    pdf_path = os.path.join(out_dir, "geological_summary_bccl.pdf")
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842) # A4
    
    text = (
        "BHARAT COKING COAL LIMITED (BCCL)\n"
        "REGIONAL GEOLOGICAL EXPLORATION REPORT\n"
        "BLOCK: JHARIA COALFIELD - SEAM XVI-A\n\n"
        "1. Geological Overview\n"
        "The Jharia coalfield represents the exclusive depository of prime coking coal in India. "
        "The Barakar Formation contains multiple seams exhibiting high thermal value and low moisture.\n\n"
        "2. Seam Correlation and Reserve Estimates\n"
        "Exploration drilling across Block 4 confirms proven reserves across the primary seams.\n\n"
    )
    page.insert_text((50, 60), text, fontsize=11, fontname="helv")
    
    # Draw table with borders so find_tables() extracts it
    x0, y0 = 50, 220
    col_widths = [100, 80, 80, 80, 100]
    row_height = 25
    headers = ["Seam Name", "Depth (m)", "Thickness (m)", "Coal Grade", "Reserve (MT)"]
    rows = [
        ["Seam XVI-A", "145.5", "4.2", "Steel-I", "12.8"],
        ["Seam XV", "210.0", "5.8", "Steel-II", "18.4"],
        ["Seam XIV", "285.2", "6.1", "Washery-I", "24.6"],
        ["Seam XIII", "340.8", "3.9", "Washery-II", "9.5"]
    ]
    
    # Draw headers
    cur_x = x0
    for i, h in enumerate(headers):
        page.draw_rect(pymupdf.Rect(cur_x, y0, cur_x + col_widths[i], y0 + row_height), color=(0, 0, 0), width=1)
        page.insert_text((cur_x + 5, y0 + 17), h, fontsize=10, fontname="helv")
        cur_x += col_widths[i]
        
    # Draw rows
    for r_idx, r in enumerate(rows):
        cur_y = y0 + (r_idx + 1) * row_height
        cur_x = x0
        for c_idx, val in enumerate(r):
            page.draw_rect(pymupdf.Rect(cur_x, cur_y, cur_x + col_widths[c_idx], cur_y + row_height), color=(0, 0, 0), width=1)
            page.insert_text((cur_x + 5, cur_y + 17), val, fontsize=10, fontname="helv")
            cur_x += col_widths[c_idx]
            
    doc.save(pdf_path)
    doc.close()
    print(f"Generated PDF: {pdf_path}")
    
    # 2. Excel Spreadsheet
    xlsx_path = os.path.join(out_dir, "monthly_coal_production_ccl.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Production_FY24"
    ws_headers = ["Month", "Subsidiary", "Area", "Target_MT", "Actual_MT", "Despatch_MT", "Compliance_Pct"]
    ws.append(ws_headers)
    ws_rows = [
        ["Apr-2023", "CCL", "Barka-Sayal", 1.25, 1.32, 1.28, 105.6],
        ["May-2023", "CCL", "Barka-Sayal", 1.30, 1.29, 1.27, 99.2],
        ["Jun-2023", "CCL", "Barka-Sayal", 1.20, 1.22, 1.21, 101.6],
        ["Jul-2023", "CCL", "Argada", 0.95, 0.98, 0.96, 103.1],
        ["Aug-2023", "CCL", "Argada", 1.05, 1.02, 1.01, 97.1],
        ["Sep-2023", "CCL", "Argada", 1.10, 1.14, 1.12, 103.6],
    ]
    for r in ws_rows:
        ws.append(r)
    wb.save(xlsx_path)
    print(f"Generated XLSX: {xlsx_path}")
    
    # 3. DOCX Inspection Report
    docx_path = os.path.join(out_dir, "mine_safety_inspection_ecl.docx")
    doc_word = docx.Document()
    doc_word.add_heading("EASTERN COALFIELDS LIMITED (ECL)", level=1)
    doc_word.add_heading("ANNUAL MINE SAFETY AUDIT REPORT - 2023-24", level=2)
    doc_word.add_paragraph(
        "This report summarizes the statutory occupational safety and DGMS compliance inspection "
        "conducted across the underground and opencast mines of the Raniganj coalfield."
    )
    doc_word.add_heading("Statutory DGMS Compliance Review", level=3)
    table = doc_word.add_table(rows=1, cols=5)
    hdr_cells = table.rows[0].cells
    doc_headers = ["Inspection Area", "Parameter", "DGMS Benchmark", "Observed Value", "Status"]
    for i, h in enumerate(doc_headers):
        hdr_cells[i].text = h
    doc_data = [
        ["Sripur Colliery", "Methane Concentration", "< 0.5%", "0.12%", "COMPLIANT"],
        ["Satgram Incline", "Air Velocity (m/min)", "> 30", "45.2", "COMPLIANT"],
        ["Jhanjra Project", "Support Resistance (t/m2)", "> 15", "18.5", "COMPLIANT"],
        ["Ningah Seam", "Dust Sampling (mg/m3)", "< 2.0", "1.4", "COMPLIANT"]
    ]
    for row in doc_data:
        row_cells = table.add_row().cells
        for i, val in enumerate(row):
            row_cells[i].text = val
    doc_word.save(docx_path)
    print(f"Generated DOCX: {docx_path}")
    
    # 4. Scanned Image (PNG)
    png_path = os.path.join(out_dir, "scanned_borehole_log.png")
    img = Image.new("RGB", (600, 300), color=(255, 255, 255))
    d = ImageDraw.Draw(img)
    d.text((30, 40), "CMPDI REGIONAL INSTITUTE-II", fill=(0, 0, 0))
    d.text((30, 80), "BOREHOLE CORE RECOVERY LOG: BH-DHN-24", fill=(0, 0, 0))
    d.text((30, 120), "Total Depth: 245.8 meters", fill=(0, 0, 0))
    d.text((30, 160), "Lithology: Sandstone with Coarse Intercalations", fill=(0, 0, 0))
    d.text((30, 200), "Logged By: Senior Geologist CMPDI RI-II", fill=(0, 0, 0))
    img.save(png_path)
    print(f"Generated PNG: {png_path}")
    
    # 5. Unsupported file (.zip)
    zip_path = os.path.join(out_dir, "unsupported_archive.zip")
    with zipfile.ZipFile(zip_path, "w") as zf:
        zf.writestr("test.txt", "unsupported dummy content")
    print(f"Generated ZIP: {zip_path}")

if __name__ == "__main__":
    generate_demo_files()
